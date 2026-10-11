"""
STEP 5 of the procedure: writing the results down.

  results/summary.csv          one row per requirement x run x mode, every number
  results/summary.xlsx         the same, formatted
  results/audit/*.md           one file per run: every matching decision behind every number
  results/answer_set_check.txt differences between the DOT/YAML encoding and Browser.pptx
"""
import csv
import os

from .labels import clean_text

COLUMNS = ["requirement", "source", "run", "mode",
           "L1_recall", "L1_precision", "L1_f1", "L1_found", "L1_answer_nodes", "L1_generated_nodes", "artifact_nodes",
           "L2_recall", "L2_precision", "L2_f1", "L2_found", "L2_answer_pairs", "L2_generated_pairs",
           "L3_recall", "L3_precision", "L3_f1", "L3_routes_walked", "L3_units", "L3_paths_needed", "L3_paths",
           "test_cases_covered", "test_cases_total", "test_case_coverage", "valid_paths", "duplicate_paths",
           "L4_recall", "L4_met", "L4_total"]


def pct(x):
    return "n/a" if x is None else f"{100 * x:.1f}"


def row_for(rid, source, run, mode, r1, r2, r3, r4):
    return {
        "requirement": rid, "source": source, "run": run, "mode": mode,
        "L1_recall": r1["recall"], "L1_precision": r1["precision"], "L1_f1": r1["f1"],
        "L1_found": r1["tp"], "L1_answer_nodes": r1["n_ref"], "L1_generated_nodes": r1["n_gen"],
        "artifact_nodes": r1["artifacts"],
        "L2_recall": r2["recall"], "L2_precision": r2["precision"], "L2_f1": r2["f1"],
        "L2_found": r2["tp"], "L2_answer_pairs": r2["n_ref"], "L2_generated_pairs": r2["n_gen"],
        "L3_recall": r3["recall"], "L3_precision": r3["precision"], "L3_f1": r3["f1"],
        "L3_routes_walked": r3["tp"], "L3_units": r3["n_ref"], "L3_paths_needed": r3["tp_paths"],
        "L3_paths": r3["n_gen"],
        "test_cases_covered": len(r3["tc_covered"]), "test_cases_total": r3["tc_total"],
        "test_case_coverage": r3["tc_coverage"], "valid_paths": r3["n_valid"], "duplicate_paths": r3["n_dup"],
        "L4_recall": r4["recall"], "L4_met": r4["met"], "L4_total": r4["n"],
    }


def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        for r in rows:
            w.writerow({k: (round(v, 4) if isinstance(v, float) else v) for k, v in r.items()})


def write_xlsx(rows, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "All results"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F4F6F")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([r[k] for k in COLUMNS])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if COLUMNS[c.column - 1].endswith(("recall", "precision", "f1", "coverage")) and isinstance(c.value, float):
                c.number_format = "0.0%"
    for i, col in enumerate(COLUMNS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(9, min(22, len(col) + 2))
    ws.freeze_panes = "E2"
    wb.save(path)
    return True


# ------------------------------------------------------------------ audit files
def _md_escape(s):
    return str(s).replace("|", "\\|").replace("\n", " ")


def write_audit(path, title, ref, spec, tg, paths, results):
    """results: {mode: (r1, r2, r3, r4)}"""
    L = []
    L.append(f"# {title}\n")
    L.append("Every number in the summary for this run comes from the tables below.\n")
    L.append("## Scores\n")
    L.append("| Level | Mode | Recall | Precision | F1 | Counted |")
    L.append("|---|---|---|---|---|---|")
    for mode, (r1, r2, r3, r4) in results.items():
        L.append(f"| L1 nodes | {mode} | {pct(r1['recall'])} | {pct(r1['precision'])} | {pct(r1['f1'])} | "
                 f"{r1['tp']} of {r1['n_ref']} answer-set nodes found; {r1['n_gen']} generated nodes |")
        L.append(f"| L2 arrows | {mode} | {pct(r2['recall'])} | {pct(r2['precision'])} | {pct(r2['f1'])} | "
                 f"{r2['tp']} of {r2['n_ref']} answer-set pairs found; {r2['n_gen']} generated pairs |")
        L.append(f"| L3 tests | {mode} | {pct(r3['recall'])} | {pct(r3['precision'])} | {pct(r3['f1'])} | "
                 f"{r3['tp']} of {r3['n_ref']} routes walked; {r3['tp_paths']} of {r3['n_gen']} paths needed; "
                 f"test cases covered {len(r3['tc_covered'])}/{r3['tc_total']} |")
        L.append(f"| L4 sparkles | {mode} | {pct(r4['recall'])} | | | {r4['met']} of {r4['n']} |")
    L.append("")

    lenient = results.get("lenient") or next(iter(results.values()))
    strict = results.get("strict") or lenient
    refl = {n.id: n.label for n in ref.concept_nodes()}

    L.append("## Step 3. What each generated node stands for\n")
    L.append("| Generated node or box | Strict | Lenient | How (lenient) | Credited at L1 (lenient) |")
    L.append("|---|---|---|---|---|")
    srows = {(r["kind"], r["id"]): r for r in strict[0]["rows"]}
    for r in lenient[0]["rows"]:
        s = srows.get((r["kind"], r["id"]))
        sm = ", ".join(refl.get(x, x) for x in s["means"]) if s else "(box not read in strict mode)"
        lm = ", ".join(refl.get(x, x) for x in r["means"]) or "nothing"
        flag = "yes" if r["credited"] else ("no, duplicate" if r["means"] else "no")
        art = " (Graphviz leftover)" if r["artifact"] else ""
        L.append(f"| {_md_escape(r['label'])}{art} | {_md_escape(sm or 'nothing')} | {_md_escape(lm)} | "
                 f"{_md_escape(r['how'])} | {flag} |")
    L.append("")
    L.append("Answer-set nodes not found (lenient): " +
             (", ".join(refl[x] for x in lenient[0]["missing"]) or "none") + "\n")

    L.append("## L2. Arrows between shared nodes (lenient)\n")
    def pairs(ps):
        return "; ".join(f"{refl.get(a, a)} -> {refl.get(b, b)}" for a, b in ps) or "none"
    L.append(f"* Found in both: {pairs(lenient[1]['matched'])}")
    L.append(f"* In the answer set only: {pairs(lenient[1]['missing'])}")
    L.append(f"* In the generated graph only: {pairs(lenient[1]['extra'])}\n")

    L.append("## L3. Test paths\n")
    L.append("How each path was read (arrow labels taken from the graph are in round brackets, "
             "arrow labels written in the path in square brackets, unknown steps between question marks):\n")
    for mode, (r1, r2, r3, r4) in results.items():
        L.append(f"### {mode}\n")
        L.append("| # | Path as read | Valid in its own graph | Routes it walks | Needed |")
        L.append("|---|---|---|---|---|")
        for p in paths:
            rts = r3["path_routes"].get(p.number, [])
            rtxt = "; ".join(" -> ".join(refl.get(c, c) for c in rt) for rt in rts) or "none"
            valid = "yes" if p.valid else "no: " + "; ".join(p.problems)
            L.append(f"| {p.number} | {_md_escape(p.pretty(tg))} | {_md_escape(valid)} | {_md_escape(rtxt)} | "
                     f"{'yes' if p.number in r3['needed'] else 'no'} |")
        L.append("")
        L.append("| Route (answer set) | Test case(s) | Walked by path(s) |")
        L.append("|---|---|---|")
        for rt in r3["routes"]:
            by = [str(n) for (n, r) in r3["hits"] if r == rt]
            L.append(f"| {_md_escape(' -> '.join(refl.get(c, c) for c in rt))} | {', '.join(r3['route_tcs'][rt])} | "
                     f"{', '.join(sorted(set(by), key=int)) or 'none'} |")
        for tc in r3["relational"]:
            L.append(f"| (compares two runs, cannot be a path) | {tc} | none |")
        L.append("")
        L.append(f"Test cases covered: {', '.join(r3['tc_covered']) or 'none'} "
                 f"({len(r3['tc_covered'])} of {r3['tc_total']}).\n")

    L.append("## L4. Sparkles\n")
    L.append("| Sparkle (from the notes) | Check | Strict | Lenient |")
    L.append("|---|---|---|---|")
    srow = {r["id"]: r for r in strict[3]["rows"]}
    for r in lenient[3]["rows"]:
        s = srow[r["id"]]
        L.append(f"| {r['id']}: {_md_escape(r['notes'])} | {r['type']} | "
                 f"{'met' if s['met'] else 'not met'}: {_md_escape(s['why'])} | "
                 f"{'met' if r['met'] else 'not met'}: {_md_escape(r['why'])} |")
    L.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(L))
